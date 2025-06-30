from flask import Flask, request, jsonify,  render_template, session
from langchain_groq import ChatGroq
from langchain_core.messages import (
    messages_from_dict, messages_to_dict
)
import json, re, os, io
from dotenv import load_dotenv
from prompts import system_prompt
from PyPDF2 import PdfReader  # For PDF parsing
import docx 
# from pdf2image import convert_from_bytes
import pytesseract
from PIL import Image
import fitz
from flask_cors import CORS
from langchain.memory import ConversationBufferMemory
from langchain_community.chat_message_histories import ChatMessageHistory
from langchain.chains import ConversationChain
from langchain.prompts import (
    ChatPromptTemplate,
    MessagesPlaceholder,
    SystemMessagePromptTemplate,
    HumanMessagePromptTemplate,
)
 
 
app =Flask(__name__)
# CORS(app)  # <-- enable CORS for all routes
CORS(app, resources={r"/*": {"origins": "*"}})


# A secret key is required for Flask to use sessions.
# In production, this should be a long, random, secret string set as an environment variable.
app.secret_key = os.getenv("FLASK_SECRET_KEY", "a-default-secret-key-for-development")

#to use API key from .env
load_dotenv()

API_KEY = os.getenv("My_API_KEY")

#ai model name
MODEL_NAME = "llama3-8b-8192"

#Langchain Setup 
llm = ChatGroq(
    model = MODEL_NAME,
    api_key=API_KEY,
    temperature=0.7,
    timeout =30
    )

# The prompt now includes placeholders for history and the new input
prompt = ChatPromptTemplate.from_messages([
    SystemMessagePromptTemplate.from_template(system_prompt["manual"]),
    MessagesPlaceholder(variable_name="chat_history"),
    HumanMessagePromptTemplate.from_template("{input}")
])

# The memory and conversation chain will be created on a per-request basis
# inside the endpoint to ensure user-specific conversation history.



#extract json from text
def extract_json_from_text(text):
    match = re.search(r'\{.*\}', text, re.DOTALL)
    if match:
        return match.group()
    return None


def extract_text_from_pdf(file_stream): #for PDF files 
    try:
        reader = PdfReader(file_stream)
        texts = []
        for pages in reader.pages:
            page_text = pages.extract_text()
            if page_text:
                texts.append(page_text)
        return '\n'.join(texts)
    except Exception as e:
        return " "

def extract_text_from_docx(file_stream):
    try:
        doc = docx.Document(file_stream)
        return '\n'.join([para.text for para in doc.paragraphs])
    except Exception as e:
        return " "

def extract_text_from_scanned_pdf_with_fitz(file_bytes):
    try:
        pdf = fitz.open(stream=file_bytes, filetype="pdf")
        extracted_text = ""
        for page in pdf:
            pix = page.get_pixmap()
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            text = pytesseract.image_to_string(img)
            extracted_text += text + "\n"
        return extracted_text.strip()
    except Exception as e:
        return ""
    
def extract_text_from_image(image_file):
    try:        
        image = Image.open(image_file)
        return pytesseract.image_to_string(image).strip()
    except Exception as e:
        return " "

#api route
@app.route('/', methods=['GET', 'POST'])
def diagnosis_prediction():
    
    if request.method == 'GET':
        #clear session history on a new get request
        session.pop('chat_history', None)
        return render_template('ChatbotUI.html')
    
    app.logger.info(f"Form Data: {request.form}")
    app.logger.info(f"Files: {request.files}")

    try:
        #POST request Logic starts here-------
        symptoms = request.form.get('symptoms', '').strip()
        age = request.form.get('age')
        gender = request.form.get('gender', '').strip().lower()
        existing_conditions = request.form.get('existing_conditions', '').strip() #optional
        report_file = request.files.get('report_file')  #optional

        if not symptoms or not age or not gender:
            return jsonify({"error": "Please provide the symptoms, age, and gender, all fields are required for better prediction"}), 400

        report_content = ""
        
        if report_file:
            file_name = report_file.filename.lower()
            
            #for handling PDF files
            if file_name.endswith('.pdf'):
                file_bytes = report_file.read()
                report_content = extract_text_from_pdf(io.BytesIO(file_bytes))
                if not report_content.strip():
                    report_content = extract_text_from_scanned_pdf_with_fitz(file_bytes)
                
            #docx handling
            elif file_name.endswith('.docx'):
                report_content = extract_text_from_docx(report_file.stream)
                
            #Image Handling
            elif file_name.endswith(('.jpg','.png','.jpeg')):
                report_content = extract_text_from_image(report_file)
            
            else:
                return jsonify({"error":"Supported file formats are .pdf, .docx, .jpg, .png, .jpeg"}), 400
                    
        
        #Construct the combined prompt for AI model
        combined_prompt = (
            f"Patient details:\n"
            f"age: {age}\n"
            f"gender: {gender}\n"
            f"symptoms: {symptoms}\n"
            f"existing conditions: {existing_conditions}\n"
        )
        
        if report_content:
            combined_prompt += f"\nReport content:\n{report_content.strip()}"
    
        #per user memory management
        #1.Load chat history from user's session
        history_dicts = session.get('chat_history', [])
        stored_messages = messages_from_dict(history_dicts)
        
        #2. Create a new memory object for this request, loading in the user's history
        memory = ConversationBufferMemory(
            chat_memory=ChatMessageHistory(messages=stored_messages),
            memory_key="chat_history",
            return_messages=True
        )
        #3. Create the conversation chain with the user-specific memory
        conversation = ConversationChain(llm = llm, prompt=prompt, memory=memory)
        
        #4. Run the prediction. The chain will automatically use the loaded history
        ai_response = conversation.predict(input=combined_prompt)
        
        #5. Save the updated conversation history backto the user's session
        updated_messages = conversation.memory.chat_memory.messages
        session['chat_history'] = messages_to_dict(updated_messages)       
        
        if ai_response:
            response_content_str = ai_response.strip()
            # app.logger.info(f"Raw AI response content:{response_content_str}")
            
            try:
                cleaned_json = extract_json_from_text(response_content_str)
                if cleaned_json:
                    parsed_json_response = json.loads(cleaned_json)
                    
                    # Check if the AI is asking a clarifying question
                    if parsed_json_response.get("status") == "clarification_needed":
                        return jsonify(parsed_json_response), 200
                    
                    # If it's not a question, it should be a full diagnosis.
                    if "diagnosis" in parsed_json_response:
                        return jsonify(parsed_json_response), 200

                # If we reach here, either no JSON was found, or the JSON was in an
                # unexpected format. Trigger the fallback.

                
                # This is a fallback. It triggers if the AI, despite the prompt,
                # returns a plain text message for irrelevant input instead of JSON.
                app.logger.warning(f"AI did not return expected JSON format. Raw Response: {response_content_str}. Creating a fallback JSON response")
                #Manully create the structured JSON response for invalid input
                fallback_response = {
                    "diagnosis": "Invalid Input",
                    "prescription": [],
                    "duration": "",
                    "tests": [],
                    "safety": "",
                    "Do and Don'ts": { "Do": [], "Don't": [] },
                    "Note": "The information provided was not recognized as a valid medical query. Please provide relevant symptoms for an accurate diagnosis."
                }
                return jsonify(fallback_response), 200
                
                
            except json.JSONDecodeError as e:
                app.logger.error(f"JSON decoding error: {e}  --Raw Response : {response_content_str}")
                return jsonify({"error": "Invalid JSON response from AI"}), 500   
        else:
            app.logger.error("AI returned an empty or invalid response")
            return jsonify({"error": "AI returned an empty or invalid response"}), 500
        
    except Exception as e:
        # catch any unexpected errors
        app.logger.error(f"Error occurred: {e}", exc_info=True)
        return jsonify({"error": "Internal server error"}), 500

if __name__ == "__main__":
    app.run(debug=True)